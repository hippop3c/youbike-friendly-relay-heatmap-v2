from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DATA_PREFIX = "window.YOUBIKE_HEATMAP_DATA="
DAILY_PREFIX = "window.YOUBIKE_DAILY_V2="
METRICS = ("滿借", "空還", "調出", "綁車", "調入", "解綁車")
CITY_BY_PREFIX = {"5001": "台北市", "5002": "新北市"}
WEEKDAY_NAMES = ("週一", "週二", "週三", "週四", "週五", "週六", "週日")


def normalize_sno(value: Any) -> str:
    text = str(value or "").strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def normalize_city(value: Any) -> str:
    return unicodedata.normalize("NFKC", str(value or "")).strip().replace("臺", "台")


def display_name(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    return re.sub(r"^YouBike2\.0[_\s]*", "", text, flags=re.I)


def normalize_name(value: Any) -> str:
    return re.sub(r"\s+", "", display_name(value).replace("臺", "台"))


def parse_datetime(value: Any) -> dt.datetime | None:
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time())
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
    ):
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return dt.datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def slot_for(value: dt.datetime) -> int:
    return value.hour * 2 + int(value.minute >= 30)


def daterange(first: dt.date, final: dt.date) -> Iterable[dt.date]:
    current = first
    while current <= final:
        yield current
        current += dt.timedelta(days=1)


def load_payload(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8").strip()
    if not text.startswith(DATA_PREFIX):
        raise RuntimeError(f"Unexpected heatmap wrapper: {path}")
    payload = text[len(DATA_PREFIX) :]
    if payload.endswith(";"):
        payload = payload[:-1]
    return json.loads(payload)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def grade_for(city: str, usage: float | None) -> str | None:
    if usage is None:
        return None
    if city == "台北市":
        return "A" if usage >= 200 else "B" if usage >= 100 else "C"
    if city == "新北市":
        return "A" if usage >= 100 else "B" if usage >= 50 else "C"
    return None


def month_configs(workspace: Path, source_root: Path) -> dict[str, dict[str, Any]]:
    return {
        "2026-07": {
            "payload": workspace / "youbike-hourly-heatmap" / "data.js",
            "start": dt.date(2026, 7, 1),
            "end": dt.date(2026, 7, 31),
            "excluded": {"2026-07-10", "2026-07-11"},
            "reward_dirs": [source_root / "raw_2026-07"],
            "vds": [
                source_root / "vds_raw_2026-07" / "vds_task_taipei_2026-07.xlsx",
                source_root / "vds_raw_2026-07" / "vds_task_newtaipei_2026-07.xlsx",
            ],
        },
        "2026-08": {
            "payload": workspace / "youbike-hourly-heatmap" / "data" / "heatmap-2026-08.js",
            "start": dt.date(2026, 8, 1),
            "end": dt.date(2026, 8, 31),
            "excluded": set(),
            "reward_dirs": [source_root / "raw_2026-08_09"],
            "vds": [
                source_root / "vds_raw_v7" / "vds_task_taipei_2026-08.xlsx",
                source_root / "vds_raw_v7" / "vds_task_newtaipei_2026-08.xlsx",
            ],
        },
        "2026-09": {
            "payload": workspace / "youbike-hourly-heatmap" / "data" / "heatmap-2026-09.js",
            "start": dt.date(2026, 9, 1),
            "end": dt.date(2026, 9, 8),
            "excluded": set(),
            "reward_dirs": [source_root / "raw_2026-08_09"],
            "vds": [
                source_root / "vds_raw_v7" / "vds_task_taipei_2026-09-01_08.xlsx",
                source_root / "vds_raw_v7" / "vds_task_newtaipei_2026-09-01_08.xlsx",
            ],
        },
    }


def build_rosters(
    configs: dict[str, dict[str, Any]], station_metadata: Path
) -> tuple[dict[str, dict[str, Any]], dict[str, defaultdict[tuple[str, str], set[int]]]]:
    rosters: dict[str, dict[str, Any]] = {}
    alias_maps: dict[str, defaultdict[tuple[str, str], set[int]]] = {}
    for month, cfg in configs.items():
        payload = load_payload(cfg["payload"])
        stations = payload["stations"]
        daily_usage = payload.get("dailyUsage") or []
        by_code: dict[str, int] = {}
        aliases: defaultdict[tuple[str, str], set[int]] = defaultdict(set)
        records = []
        for index, row in enumerate(stations):
            code = normalize_sno(row[5] if len(row) > 5 else "")
            city = normalize_city(row[1])
            usage_raw = daily_usage[index] if index < len(daily_usage) else None
            usage = float(usage_raw) if isinstance(usage_raw, (int, float)) else None
            record = {
                "index": index,
                "code": code,
                "name": display_name(row[0]),
                "city": city,
                "district": str(row[2] or "").strip(),
                "usage": usage,
                "grade": grade_for(city, usage),
            }
            records.append(record)
            if code:
                by_code[code] = index
            aliases[(city, normalize_name(row[0]))].add(index)
        rosters[month] = {"payload": payload, "records": records, "by_code": by_code}
        alias_maps[month] = aliases

    metadata = json.loads(station_metadata.read_text(encoding="utf-8"))
    for month, roster in rosters.items():
        aliases = alias_maps[month]
        for item in metadata.get("stations", []):
            city = normalize_city(item.get("city"))
            code = normalize_sno(item.get("stationId"))
            names = [item.get("stationName"), *(item.get("aliases") or [])]
            target = roster["by_code"].get(code)
            if target is None:
                candidates = {
                    index
                    for name in names
                    if name
                    for index in aliases.get((city, normalize_name(name)), set())
                }
                if len(candidates) == 1:
                    target = next(iter(candidates))
            if target is None:
                continue
            if code:
                roster["by_code"].setdefault(code, target)
            for name in names:
                if name:
                    aliases[(city, normalize_name(name))].add(target)
    return rosters, alias_maps


def resolve_station(
    month: str,
    city_value: Any,
    code_value: Any,
    name_value: Any,
    rosters: dict[str, dict[str, Any]],
    alias_maps: dict[str, defaultdict[tuple[str, str], set[int]]],
    audit: Counter,
) -> int | None:
    city = normalize_city(city_value)
    code = normalize_sno(code_value)
    roster = rosters[month]
    index = roster["by_code"].get(code)
    if index is not None and roster["records"][index]["city"] == city:
        return index
    candidates = sorted(alias_maps[month].get((city, normalize_name(name_value)), set()))
    if candidates:
        if len(candidates) > 1:
            audit["ambiguousStationMatches"] += 1
        return candidates[0]
    audit["unmatchedStations"] += 1
    return None


def reward_files(configs: dict[str, dict[str, Any]]) -> list[Path]:
    unique: dict[str, Path] = {}
    for cfg in configs.values():
        for directory in cfg["reward_dirs"]:
            for path in sorted(directory.glob("*.csv")):
                unique[str(path.resolve()).lower()] = path
    return list(unique.values())


def load_rewards(
    configs: dict[str, dict[str, Any]],
    rosters: dict[str, dict[str, Any]],
    alias_maps: dict[str, defaultdict[tuple[str, str], set[int]]],
    cells: dict[str, dict[tuple[str, int, int], list[int]]],
) -> dict[str, Any]:
    fields = ("帳號", "分類", "借車時間", "借車縣市", "借車場站", "還車時間", "還車縣市", "還車場站")
    orders: dict[str, list[str]] = {}
    audit = Counter()
    sources = []
    for path in reward_files(configs):
        file_rows = 0
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not {"訂單號", *fields}.issubset(reader.fieldnames or []):
                continue
            for row in reader:
                file_rows += 1
                audit["sourceRows"] += 1
                order_id = str(row.get("訂單號") or "").strip()
                if not order_id:
                    audit["missingOrderRows"] += 1
                    continue
                values = [str(row.get(field) or "").strip() for field in fields]
                existing = orders.get(order_id)
                if existing is None:
                    orders[order_id] = values
                else:
                    audit["duplicateRows"] += 1
                    existing[1] = " ".join(sorted(set(existing[1].split()) | set(values[1].split())))
                    for index, value in enumerate(values):
                        if not existing[index] and value:
                            existing[index] = value
        sources.append({"file": path.name, "rows": file_rows, "bytes": path.stat().st_size})
    audit["uniqueOrders"] = len(orders)

    event_specs = (("滿借", 0, 2, 3, 4), ("空還", 1, 5, 6, 7))
    for values in orders.values():
        if not values[0].startswith("09"):
            audit["excludedNonMobileOrders"] += 1
            continue
        audit["mobileOrders"] += 1
        for label, metric_index, time_index, city_index, name_index in event_specs:
            if label not in values[1]:
                continue
            timestamp = parse_datetime(values[time_index])
            if timestamp is None:
                audit["badTimestamps"] += 1
                continue
            month = timestamp.strftime("%Y-%m")
            cfg = configs.get(month)
            if cfg is None or not cfg["start"] <= timestamp.date() <= cfg["end"]:
                continue
            date_text = timestamp.date().isoformat()
            if date_text in cfg["excluded"]:
                audit[f"excluded{label}"] += 1
                continue
            city = normalize_city(values[city_index])
            station_index = resolve_station(
                month, city, "", values[name_index], rosters, alias_maps, audit
            )
            if station_index is None:
                continue
            key = (date_text, slot_for(timestamp), station_index)
            target = cells[month].setdefault(key, [0, 0, 0, 0, 0, 0])
            target[metric_index] += 1
            audit[f"included{label}"] += 1
    return {"counts": dict(audit), "files": sources}


def load_vds(
    configs: dict[str, dict[str, Any]],
    rosters: dict[str, dict[str, Any]],
    alias_maps: dict[str, defaultdict[tuple[str, str], set[int]]],
    cells: dict[str, dict[tuple[str, int, int], list[int]]],
) -> dict[str, Any]:
    action_index = {"調出": 2, "綁車": 3, "調入": 4, "解車": 5, "解綁車": 5}
    audit = Counter()
    sources = []
    for month, cfg in configs.items():
        for path in cfg["vds"]:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator)]
            positions = {name: index for index, name in enumerate(headers)}
            is_new = {"工作狀態", "數量", "更新時間"}.issubset(positions)
            is_old = {"實際調度", "任務狀態", "更新時間"}.issubset(positions)
            if not is_new and not is_old:
                workbook.close()
                raise RuntimeError(f"Unsupported VDS schema: {path.name}")
            file_counts = Counter()
            for raw in iterator:
                file_counts["sourceRows"] += 1
                audit["sourceRows"] += 1
                if is_old and str(raw[positions["任務狀態"]] or "").strip() != "完成":
                    continue
                timestamp = parse_datetime(raw[positions["更新時間"]])
                if timestamp is None or not cfg["start"] <= timestamp.date() <= cfg["end"]:
                    continue
                if timestamp.date().isoformat() in cfg["excluded"]:
                    audit["excludedDateRows"] += 1
                    continue
                if is_new:
                    action_text = str(raw[positions["工作狀態"]] or "").strip()
                    metric_index = action_index.get(action_text)
                    if metric_index is None:
                        continue
                    try:
                        quantity = int(float(raw[positions["數量"]] or 0))
                    except (TypeError, ValueError):
                        audit["invalidQuantity"] += 1
                        continue
                else:
                    action_text = str(raw[positions["實際調度"]] or "").strip()
                    match = re.search(r"^(調入|調出|綁車|解車|解綁車)\s*(\d+)", action_text)
                    if not match:
                        continue
                    metric_index = action_index[match.group(1)]
                    action_text = match.group(1)
                    quantity = int(match.group(2))
                if quantity <= 0:
                    continue
                code = normalize_sno(raw[positions.get("場站代號", 2)])
                city = normalize_city(raw[positions.get("城市", 0)]) or CITY_BY_PREFIX.get(code[:4], "")
                name = raw[positions.get("場站名稱", 3)]
                station_index = resolve_station(
                    month, city, code, name, rosters, alias_maps, audit
                )
                if station_index is None:
                    continue
                key = (timestamp.date().isoformat(), slot_for(timestamp), station_index)
                target = cells[month].setdefault(key, [0, 0, 0, 0, 0, 0])
                target[metric_index] = max(target[metric_index], quantity)
                metric_name = METRICS[metric_index]
                audit[f"included{metric_name}"] += 1
                file_counts[metric_name] += 1
            workbook.close()
            sources.append(
                {
                    "file": path.name,
                    "month": month,
                    "schema": "staff-info" if is_new else "task-summary",
                    "counts": dict(file_counts),
                    "bytes": path.stat().st_size,
                }
            )
    return {"counts": dict(audit), "files": sources}


def serialize_month(
    month: str,
    cfg: dict[str, Any],
    roster: dict[str, Any],
    cells: dict[tuple[str, int, int], list[int]],
    output_dir: Path,
    audit: dict[str, Any],
) -> dict[str, Any]:
    dates = [value for value in daterange(cfg["start"], cfg["end"]) if value.isoformat() not in cfg["excluded"]]
    date_index = {value.isoformat(): index for index, value in enumerate(dates)}
    daily: list[list[list[list[int]]]] = [[[] for _ in range(48)] for _ in dates]
    for (date_text, slot, station_index), values in sorted(cells.items()):
        target_date = date_index.get(date_text)
        if target_date is None or not any(values):
            continue
        daily[target_date][slot].append([station_index, *values])

    payload = roster["payload"]
    output = {
        "version": "v2",
        "month": month,
        "metrics": list(METRICS),
        "dates": [
            {
                "date": value.isoformat(),
                "dayType": "平日" if value.weekday() < 5 else "假日",
                "weekday": WEEKDAY_NAMES[value.weekday()],
            }
            for value in dates
        ],
        "dailyActivity": daily,
        "meta": {
            "periodStart": cfg["start"].isoformat(),
            "periodEnd": cfg["end"].isoformat(),
            "excludedDates": sorted(cfg["excluded"]),
            "stationCount": len(roster["records"]),
            "dateCount": len(dates),
            "weekdayCount": sum(value.weekday() < 5 for value in dates),
            "holidayCount": sum(value.weekday() >= 5 for value in dates),
            "timeResolutionMinutes": 30,
            "eventAggregation": "unique-mobile-order-count-per-station-date-half-hour",
            "dispatchAggregation": "maximum-task-vehicle-count-per-station-date-half-hour-action",
            "multiDateAggregation": "arithmetic-mean-including-zero-days",
            "gradeThresholds": {
                "台北市": {"A": ">=200", "B": "100-199.99", "C": "<100"},
                "新北市": {"A": ">=100", "B": "50-99.99", "C": "<50"},
            },
            "sourceMonthlyPayloadSha256": file_sha256(cfg["payload"]),
            "monthlyPayloadPeriod": payload.get("meta", {}).get("period"),
            "audit": audit,
        },
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"daily-{month}.js"
    output_path.write_text(
        DAILY_PREFIX + json.dumps(output, ensure_ascii=False, separators=(",", ":")) + ";",
        encoding="utf-8",
    )
    return {
        "month": month,
        "output": str(output_path),
        "bytes": output_path.stat().st_size,
        "sha256": file_sha256(output_path),
        "dates": len(dates),
        "cells": len(cells),
        "rows": sum(len(bucket) for day in daily for bucket in day),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workspace",
        type=Path,
        default=Path(__file__).resolve().parents[2],
        help="Workspace root containing youbike-hourly-heatmap",
    )
    parser.add_argument(
        "--source-root",
        type=Path,
        default=Path.home() / "Documents" / "暫停營運測試" / "cps_reward_log",
    )
    parser.add_argument(
        "--station-metadata",
        type=Path,
        default=Path(__file__).resolve().parents[2]
        / "outputs"
        / "01a057bd-84e6-7bd3-995f-1d34ab3e518e"
        / "hourly_metrics_data"
        / "stations.json",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "dist" / "data",
    )
    parser.add_argument(
        "--summary",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "build-summary.json",
    )
    args = parser.parse_args()

    configs = month_configs(args.workspace, args.source_root)
    required = [args.station_metadata]
    for cfg in configs.values():
        required.extend([cfg["payload"], *cfg["reward_dirs"], *cfg["vds"]])
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise RuntimeError("Missing sources: " + json.dumps(missing, ensure_ascii=False))

    rosters, alias_maps = build_rosters(configs, args.station_metadata)
    cells: dict[str, dict[tuple[str, int, int], list[int]]] = {
        month: {} for month in configs
    }
    reward_audit = load_rewards(configs, rosters, alias_maps, cells)
    vds_audit = load_vds(configs, rosters, alias_maps, cells)

    outputs = []
    for month, cfg in configs.items():
        outputs.append(
            serialize_month(
                month,
                cfg,
                rosters[month],
                cells[month],
                args.output_dir,
                {"reward": reward_audit["counts"], "vds": vds_audit["counts"]},
            )
        )
    summary = {
        "version": "v2",
        "generatedAt": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "outputs": outputs,
        "rewardSources": reward_audit["files"],
        "vdsSources": vds_audit["files"],
    }
    args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
